# backend/routes/afterschool_synopsis.py
"""
After-School / ECA Synopsis routes.

Separate, additive feature from the existing camp-synopsis feature
(routes/synopsis.py) — new Firestore collections (`afterschool_synopsis`,
`afterschool_synopsis_months`), new route prefix. Reuses (never modifies)
routes/synopsis.py's ICC-admin auth dependency and .docx-building primitives,
per the approved implementation plan.

Two audiences:
- Teachers: unauthenticated, gated by a single shared portal credential
  (AfterschoolSynopsisConfig.PORTAL_USERNAME/PASSWORD) checked client-side
  after POST /login — not a per-request auth dependency, matching the old
  camp feature's teacher-facing entries/photos endpoints.
- Admins: Firebase-token gated via verify_icc_admin (@indiacc.org), imported
  directly from routes/synopsis.py.
"""

import hmac
import io
import logging
import os as _os
import uuid
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as _qn

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import Response

from config import AfterschoolSynopsisConfig
from schemas.afterschool_synopsis_schema import (
    AFTER_SCHOOL_CLASS_TYPE,
    COLOR_THEME_PALETTE,
    ECA_TYPE_ORDER,
    GRADE_SLUG_TO_LABEL,
    GRADES_WITHOUT_ECA,
    PHOTO_MAX,
    SYNOPSIS_TYPE_OPTIONS,
    AfterschoolBlockFields as BF,
    AfterschoolEntryFields as EF,
    AfterschoolMonthFields as MF,
    BlockInput,
    EnhanceTextRequest,
    EntrySaveRequest,
    MonthCreate,
    MonthUpdate,
    PortalLoginRequest,
    default_theme_for_month,
    is_allowed_month,
    month_id_for,
    month_label_for,
    slugify,
    synopsis_types_for_grade,
)
from services.firebase_service import FirebaseService
from utils.llm_handler import call_openai, OpenAIServiceError

# Reuse (never duplicate) the camp-synopsis feature's admin-auth dependency and
# .docx-building primitives — per the approved plan, this is genuine code
# reuse, not reimplementation. routes/synopsis.py itself is never modified.
from routes.synopsis import (
    verify_icc_admin,
    _run,
    _para,
    _para_border_bottom,
    _cell_bg,          # noqa: F401 — imported for parity/future use, not currently called
    _remove_table_borders,
    _add_hyperlink,
    _add_md_text,
    _fetch_photo_bytes,
    _fix_exif_rotation,
    _heic_to_jpeg,
    _LOGO_PATH,
    ALLOWED_IMAGE_TYPES,
    HEIC_IMAGE_TYPES,
    MAX_PHOTO_BYTES,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/afterschool-synopsis", tags=["afterschool_synopsis"])
firebase = FirebaseService()

# Deliberately different from the summer camp ENHANCE_SYSTEM_PROMPT in
# synopsis.py: after-school synopses showcase everything students covered, so
# this prompt must be comprehensive and structured, never condensed. Output
# formatting is limited to **bold**, "•" bullets and line breaks because that is
# all the textarea and the DOCX export (_add_md_text) render cleanly.
AFTERSCHOOL_ENHANCE_SYSTEM_PROMPT = (
    "You help teachers turn their notes into polished, parent-facing summaries for a K-6 "
    "after-school and enrichment program. The goal is to showcase EVERYTHING the students "
    "learned and did, so the result must be comprehensive, not a condensed summary.\n\n"
    "Completeness (most important):\n"
    "- Keep every single piece of information the teacher wrote: every concept, idea, skill, "
    "activity, project, tool, material, book, game, name, number, date, and outcome.\n"
    "- Never drop, merge away, or generalize a detail to make the text shorter. If the teacher "
    "listed five items, all five must appear. Length is not a concern.\n"
    "- Do not add, invent, or elaborate on anything the teacher did not write.\n\n"
    "Style:\n"
    "- Paraphrase only to make the writing professional, clean, and easy to read, and fix "
    "spelling and grammar.\n"
    "- Use a warm, first-person plural voice (e.g. 'we practiced...', 'our students...').\n"
    "- Do not use em dashes (—). Use commas or short sentences instead.\n\n"
    "Structure:\n"
    "- Open with one short sentence introducing what the class explored.\n"
    "- Then group the content into clear sections that fit the material, such as "
    "Concepts & Ideas, Skills Practiced, Activities & Projects (only use sections the "
    "teacher's note actually supports).\n"
    "- Put each section label on its own line in bold with a fitting emoji, e.g. "
    "'🧠 **Concepts & Ideas**'.\n"
    "- Under each label, list the items as bullet lines starting with '• ', one item per line, "
    "each a short, complete phrase or sentence.\n"
    "- Leave a blank line between sections.\n"
    "- If the note is very short and covers only one idea, a single well-written paragraph "
    "is fine instead of sections.\n"
    "- Use relevant emojis on section labels and where they naturally fit; do not overdo it.\n"
    "- Formatting is limited to **bold**, '• ' bullets, and line breaks. Do not use '#' "
    "headings, '-' or '*' list markers, numbered lists, or tables.\n\n"
    "Return only the finished text, with no preamble and no quotes around it."
)


# ── Teacher-facing: portal login ─────────────────────────────────────────────

@router.post("/login")
async def portal_login(body: PortalLoginRequest):
    """Shared-credential check, not per-teacher auth. Constant-time compare
    to avoid leaking timing information about the stored credential."""
    user_ok = hmac.compare_digest(body.username, AfterschoolSynopsisConfig.PORTAL_USERNAME)
    pass_ok = hmac.compare_digest(body.password, AfterschoolSynopsisConfig.PORTAL_PASSWORD)
    if not (user_ok and pass_ok):
        raise HTTPException(status_code=401, detail="Invalid portal credentials")
    return {"success": True}


# ── Teacher-facing: months ───────────────────────────────────────────────────

@router.get("/months/visible")
async def list_visible_months():
    months = await firebase.list_visible_months()
    return {"months": months}


@router.get("/months/active")
async def get_active_month_route():
    month = await firebase.get_active_month()
    return {"month": month}


# ── Teacher-facing: entries ──────────────────────────────────────────────────

@router.get("/entries/{grade_slug}/{type_slug}/{month_id}")
async def get_entry(grade_slug: str, type_slug: str, month_id: str):
    """Direct doc lookup — no query, no composite index required."""
    entry_id = f"{grade_slug}__{type_slug}__{month_id}"
    entry = await firebase.get_afterschool_entry(entry_id)
    return {"entry": entry}


@router.post("/entries")
async def save_entry(body: EntrySaveRequest):
    if body.synopsis_type not in SYNOPSIS_TYPE_OPTIONS:
        raise HTTPException(400, f"Invalid synopsis_type. Must be one of: {', '.join(SYNOPSIS_TYPE_OPTIONS)}")

    grade_slug = slugify(body.grade)
    if grade_slug not in GRADE_SLUG_TO_LABEL:
        raise HTTPException(400, f"Invalid grade '{body.grade}'")
    if body.synopsis_type not in synopsis_types_for_grade(GRADE_SLUG_TO_LABEL[grade_slug]):
        raise HTTPException(400, f"{GRADE_SLUG_TO_LABEL[grade_slug]} has no {body.synopsis_type} synopsis")
    type_slug = slugify(body.synopsis_type)

    blocks = body.blocks

    # Every synopsis_type is single-block now (SINGLE_BLOCK_TYPES == all types).
    if len(blocks) != 1:
        raise HTTPException(400, "Exactly 1 entry block is required")
    b = blocks[0]
    blocks = [BlockInput(week=None, title=b.title, raw_text=b.raw_text, photo_urls=b.photo_urls)]

    for b in blocks:
        if len(b.photo_urls) > PHOTO_MAX:
            raise HTTPException(400, f"Each block may have at most {PHOTO_MAX} photos")

    month = await firebase.get_month(body.month_id)
    month_label = month.get(MF.LABEL) if month else body.month_id

    entry_id = f"{grade_slug}__{type_slug}__{body.month_id}"
    now = datetime.utcnow().isoformat()
    existing = await firebase.get_afterschool_entry(entry_id)

    data = {
        EF.ENTRY_ID: entry_id,
        EF.GRADE: GRADE_SLUG_TO_LABEL[grade_slug],
        EF.GRADE_SLUG: grade_slug,
        EF.SYNOPSIS_TYPE: body.synopsis_type,
        EF.TYPE_SLUG: type_slug,
        EF.MONTH_ID: body.month_id,
        EF.MONTH_LABEL: month_label,
        EF.BLOCKS: [b.model_dump() for b in blocks],
        EF.DRIVE_LINK: (body.drive_link or ''),
        EF.CREATED_AT: (existing or {}).get(EF.CREATED_AT, now),
        EF.UPDATED_AT: now,
    }

    await firebase.upsert_afterschool_entry(entry_id, data)
    return {"success": True, "entry_id": entry_id}


@router.post("/enhance")
async def enhance_text(body: EnhanceTextRequest):
    """Enhance a single block's description with AI — teacher-triggered only."""
    if not body.raw_text.strip():
        raise HTTPException(400, "raw_text is required")
    try:
        result = call_openai(
            prompt=body.raw_text,
            system_message=AFTERSCHOOL_ENHANCE_SYSTEM_PROMPT,
            json_mode=False,
            temperature=0.3,
        )
    except OpenAIServiceError as e:
        raise HTTPException(status_code=503, detail=str(e))
    return {"enhanced_text": result.get("response", "")}


# ── Teacher-facing: photos ───────────────────────────────────────────────────

@router.post("/photos")
async def upload_photo(
    file: UploadFile = File(...),
    grade_slug: str = Query(...),
    type_slug: str = Query(...),
    month_id: str = Query(...),
    block_index: int = Query(...),
):
    """Mirrors routes/synopsis.py's POST /api/synopsis/photos byte-for-byte in
    approach: same allowed types, same HEIC transcoding, same size cap, same
    upload_file() storage helper — just a different storage path prefix."""
    content_type = file.content_type or ""
    is_heic = content_type in HEIC_IMAGE_TYPES
    if not is_heic and content_type not in ALLOWED_IMAGE_TYPES:
        raise HTTPException(400, "Only JPEG, PNG, WebP, GIF, and HEIC images are supported")

    data = await file.read()
    if len(data) > MAX_PHOTO_BYTES:
        raise HTTPException(400, "Image must be under 10 MB")

    ext = file.filename.rsplit(".", 1)[-1] if file.filename and "." in file.filename else "jpg"

    if is_heic:
        try:
            data = await run_in_threadpool(_heic_to_jpeg, data)
        except Exception as exc:
            logger.warning(f"HEIC conversion failed for {file.filename}: {exc}")
            raise HTTPException(400, "Couldn't process this HEIC photo. Please try converting it to JPEG first.")
        content_type = "image/jpeg"
        ext = "jpg"

    photo_id = str(uuid.uuid4())
    storage_path = f"afterschool_synopsis/{grade_slug}/{type_slug}/{month_id}/{block_index}/{photo_id}.{ext}"

    url = await firebase.upload_file(data, storage_path, content_type)
    return {"url": url}


# ── Admin-facing: months ─────────────────────────────────────────────────────

@router.post("/months", status_code=status.HTTP_201_CREATED)
async def create_month(body: MonthCreate, admin: dict = Depends(verify_icc_admin)):
    if not is_allowed_month(body.year, body.month):
        raise HTTPException(400, "Month must be one of September 2026 through May 2027")
    month_id = month_id_for(body.year, body.month)
    if await firebase.get_month(month_id):
        raise HTTPException(409, f"Month {month_id} already exists")

    now = datetime.utcnow().isoformat()
    data = {
        MF.MONTH_ID: month_id,
        MF.YEAR: body.year,
        MF.MONTH: body.month,
        MF.LABEL: month_label_for(body.year, body.month),
        MF.IS_ACTIVE: False,
        MF.IS_VISIBLE: True,
        MF.COLOR_THEME: default_theme_for_month(body.month),
        MF.CREATED_BY: admin.get("email", admin.get("uid", "")),
        MF.CREATED_AT: now,
    }
    await firebase.create_month(month_id, data)
    return {"month_id": month_id}


@router.patch("/months/{month_id}")
async def update_month(month_id: str, body: MonthUpdate, admin: dict = Depends(verify_icc_admin)):
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    if not updates:
        raise HTTPException(400, "No fields to update")
    if MF.COLOR_THEME in updates and updates[MF.COLOR_THEME] not in COLOR_THEME_PALETTE:
        raise HTTPException(400, f"color_theme must be one of: {', '.join(COLOR_THEME_PALETTE)}")
    if updates.get(MF.IS_ACTIVE):
        await firebase.deactivate_all_afterschool_months()
    ok = await firebase.update_month(month_id, updates)
    if not ok:
        raise HTTPException(404, "Month not found")
    return {"success": True}


@router.delete("/months/{month_id}")
async def delete_month(month_id: str, admin: dict = Depends(verify_icc_admin)):
    ok = await firebase.delete_month(month_id)
    if not ok:
        raise HTTPException(404, "Month not found")
    return {"success": True}


@router.get("/months")
async def list_all_months(admin: dict = Depends(verify_icc_admin)):
    months = await firebase.list_months()
    return {"months": months}


# ── Admin-facing: class view ─────────────────────────────────────────────────

def _entry_key(synopsis_type: str) -> str:
    """'After School Class' -> 'after_school_class', 'Theater' -> 'theater', etc."""
    return slugify(synopsis_type).replace('-', '_')


@router.get("/classes/{grade_slug}")
async def get_class_entries(
    grade_slug: str,
    month_id: str = Query(...),
    admin: dict = Depends(verify_icc_admin),
):
    if grade_slug not in GRADE_SLUG_TO_LABEL:
        raise HTTPException(404, "Unknown grade")
    grade_label = GRADE_SLUG_TO_LABEL[grade_slug]
    month = await firebase.get_month(month_id)

    entries = {}
    for synopsis_type in synopsis_types_for_grade(grade_label):
        type_slug = slugify(synopsis_type)
        entry_id = f"{grade_slug}__{type_slug}__{month_id}"
        entries[_entry_key(synopsis_type)] = await firebase.get_afterschool_entry(entry_id)

    return {"grade": grade_label, "month": month, "entries": entries}


# ── Newsletter document design ───────────────────────────────────────────────
# Reuses _run/_para/_para_border_bottom/_remove_table_borders/_add_md_text and
# the real ICC logo directly from routes/synopsis.py, mirroring
# _build_synopsis_doc()'s structure exactly per the approved plan — no new
# doc-building primitives invented here.

_WEEK_COLOR_EMOJI = [
    ('B8E8A5', '🌟'),
    ('A5C9E8', '⭐'),
    ('E8D5A5', '🌈'),
    ('E8A5A5', '🎉'),
    ('C5B8E8', '🎊'),
]


def _resolve_band_color(month: dict) -> str:
    theme = (month or {}).get(MF.COLOR_THEME)
    return COLOR_THEME_PALETTE.get(theme, COLOR_THEME_PALETTE['yellow'])


def _school_year_label(year: int, month: int) -> str:
    """e.g. (2026, 9) -> '2026-27' — a school-year range, not a single
    calendar year, and derived from the selected month, not generation date."""
    start_year = year if month >= 8 else year - 1
    return f"{start_year}-{str((start_year + 1) % 100).zfill(2)}"


def _add_header_footer_title_block(doc: Document, *, grade_label: str, meta_text: str, school_year_label: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.1)
    section.bottom_margin = Inches(1.1)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    for hf_p in (section.header.paragraphs[0], section.footer.paragraphs[0]):
        hf_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hf_r = hf_p.add_run(f'ICC After School Program {school_year_label}')
        hf_r.font.size = Pt(11)
        hf_r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    ttbl = doc.add_table(rows=1, cols=2)
    ttbl.autofit = False
    ttbl.columns[0].width = Inches(4.5)
    ttbl.columns[1].width = Inches(2.0)
    _remove_table_borders(ttbl)

    cl = ttbl.cell(0, 0)
    p_name = cl.paragraphs[0]
    p_name.paragraph_format.space_before = Pt(6)
    p_name.paragraph_format.space_after = Pt(6)
    r_title = p_name.add_run(grade_label)
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    p_meta = cl.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(4)
    p_meta.paragraph_format.space_after = Pt(10)
    r_meta = p_meta.add_run(meta_text)
    r_meta.font.size = Pt(12)
    r_meta.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    cr = ttbl.cell(0, 1)
    p_logo = cr.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if _os.path.exists(_LOGO_PATH):
        try:
            p_logo.add_run().add_picture(_LOGO_PATH, width=Inches(1.9))
        except Exception:
            pass

    p_div = _para(doc, spc_b=6, spc_a=10)
    _para_border_bottom(p_div, '1a1a1a', sz='12')


def _add_section_heading(doc: Document, text: str, band_hex: str, page_break: bool) -> None:
    p = _para(doc, spc_b=16, spc_a=10, align=WD_ALIGN_PARAGRAPH.CENTER, para_bg=band_hex)
    if page_break:
        p.paragraph_format.page_break_before = True
    _run(p, text, bold=True, size_pt=23)
    _para_border_bottom(p, '1a1a1a', sz='6')


def _add_drive_link_bar(doc: Document, label: str, url: str) -> None:
    """One Google Drive link bar, placed directly under its section heading.
    Same A5C9E8 shading / 📷 convention as routes/synopsis.py's gallery bar,
    but the "Google Drive link" label is set large and bold on its own line so
    it stands out in the doc; the URL itself stays at 10pt below it."""
    p_link = _para(doc, spc_b=4, spc_a=8)
    pPr = p_link._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(_qn('w:val'), 'clear')
    shd.set(_qn('w:color'), 'auto')
    shd.set(_qn('w:fill'), 'A5C9E8')
    pPr.append(shd)
    _run(p_link, f'📷  Google Drive link for {label}', bold=True, size_pt=15)
    p_link.add_run().add_break()
    _add_hyperlink(p_link, url, url, size_pt=10, bold=True)


def _add_intro(doc: Document, title: str, text: str) -> None:
    """Admin-written monthly intro paragraph — sits between the title/logo
    block and the first section heading of the After School newsletter."""
    if title:
        p_title = _para(doc, spc_b=4, spc_a=6)
        _run(p_title, title, bold=True, size_pt=16)
    if text:
        p_text = _para(doc, spc_b=2, spc_a=12, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _add_md_text(p_text, text, size_pt=11)


def _add_block_entry(doc: Document, block: dict, *, week_index: int, label_text: str) -> None:
    color_hex, emoji = _WEEK_COLOR_EMOJI[week_index % 5]
    p_label = _para(doc, spc_b=12, spc_a=6)
    _run(p_label, f'{emoji} {label_text}', bold=True, size_pt=14, bg_hex=color_hex)

    raw_text = (block or {}).get(BF.RAW_TEXT) or ''
    p_notes = _para(doc, spc_b=6, spc_a=12, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if raw_text:
        _add_md_text(p_notes, raw_text, size_pt=11)
    else:
        _run(p_notes, 'No notes submitted for this month.', size_pt=11)

    photo_urls = (block or {}).get(BF.PHOTO_URLS) or []
    if photo_urls:
        for row_start in range(0, len(photo_urls), 3):
            p_photos = _para(doc, spc_b=4, spc_a=4, align=WD_ALIGN_PARAGRAPH.CENTER)
            for url in photo_urls[row_start:row_start + 3]:
                try:
                    img_b = _fix_exif_rotation(_fetch_photo_bytes(url))
                    p_photos.add_run().add_picture(io.BytesIO(img_b), width=Inches(2.1))
                except Exception as exc:
                    logger.warning(f"Could not embed photo {url}: {exc}")

    p_sep = _para(doc, spc_b=2, spc_a=2)
    _para_border_bottom(p_sep, 'e5e5e5', sz='2')


def _build_after_school_doc(*, grade_label: str, month: dict, entry: Optional[dict]) -> bytes:
    """After School Class newsletter — one section, one entry block."""
    band_hex = _resolve_band_color(month)
    month_label = month.get(MF.LABEL, '')
    school_year_label = _school_year_label(month.get(MF.YEAR), month.get(MF.MONTH))

    doc = Document()
    _add_header_footer_title_block(
        doc,
        grade_label=grade_label,
        meta_text=f'After School Class · {month_label}',
        school_year_label=school_year_label,
    )

    intro_title = (month.get(MF.INTRO_TITLE) or '').strip()
    intro_text = (month.get(MF.INTRO_TEXT) or '').strip()
    if intro_title or intro_text:
        _add_intro(doc, intro_title, intro_text)

    _add_section_heading(doc, AFTER_SCHOOL_CLASS_TYPE, band_hex, page_break=False)

    drive_link = (entry or {}).get(EF.DRIVE_LINK)
    if drive_link:
        _add_drive_link_bar(doc, 'Class Photos', drive_link)

    blocks = (entry or {}).get(EF.BLOCKS) or []
    if blocks:
        block = blocks[0]
        label_text = block.get(BF.TITLE) or AFTER_SCHOOL_CLASS_TYPE
        _add_block_entry(doc, block, week_index=0, label_text=label_text)
    else:
        p_notes = _para(doc, spc_b=6, spc_a=12, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _run(p_notes, 'No notes submitted for this month.', size_pt=11)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _build_eca_doc(*, grade_label: str, month: dict, entries_by_type: dict) -> bytes:
    """ECA newsletter — one section per ECA (fixed order), page break before
    each after the first; one title+description+photos block per ECA, same
    shape as _build_after_school_doc()'s single-block treatment."""
    band_hex = _resolve_band_color(month)
    month_label = month.get(MF.LABEL, '')
    school_year_label = _school_year_label(month.get(MF.YEAR), month.get(MF.MONTH))

    doc = Document()
    _add_header_footer_title_block(
        doc,
        grade_label=grade_label,
        meta_text=f'ECA Newsletter · {month_label}',
        school_year_label=school_year_label,
    )

    for idx, eca_type in enumerate(ECA_TYPE_ORDER):
        entry = entries_by_type.get(eca_type)
        _add_section_heading(doc, eca_type, band_hex, page_break=(idx > 0))

        # Each ECA's Google Drive link sits directly under its own heading.
        drive_link = (entry or {}).get(EF.DRIVE_LINK)
        if drive_link:
            _add_drive_link_bar(doc, eca_type, drive_link)

        blocks = (entry or {}).get(EF.BLOCKS) or []
        if blocks:
            block = blocks[0]
            label_text = block.get(BF.TITLE) or eca_type
            _add_block_entry(doc, block, week_index=0, label_text=label_text)
        else:
            p_notes = _para(doc, spc_b=6, spc_a=12, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            _run(p_notes, 'No notes submitted for this month.', size_pt=11)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@router.get("/classes/{grade_slug}/download/after-school")
async def download_after_school_doc(
    grade_slug: str,
    month_id: str = Query(...),
    admin: dict = Depends(verify_icc_admin),
):
    if grade_slug not in GRADE_SLUG_TO_LABEL:
        raise HTTPException(404, "Unknown grade")
    grade_label = GRADE_SLUG_TO_LABEL[grade_slug]
    month = await firebase.get_month(month_id)
    if not month:
        raise HTTPException(404, "Month not found")

    type_slug = slugify(AFTER_SCHOOL_CLASS_TYPE)
    entry_id = f"{grade_slug}__{type_slug}__{month_id}"
    entry = await firebase.get_afterschool_entry(entry_id)

    doc_bytes = await run_in_threadpool(
        _build_after_school_doc, grade_label=grade_label, month=month, entry=entry
    )
    slug = slugify(grade_label)
    filename = f'after_school_{slug}_{month_id}.docx'
    return Response(
        content=doc_bytes,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )


@router.get("/classes/{grade_slug}/download/eca")
async def download_eca_doc(
    grade_slug: str,
    month_id: str = Query(...),
    admin: dict = Depends(verify_icc_admin),
):
    if grade_slug not in GRADE_SLUG_TO_LABEL:
        raise HTTPException(404, "Unknown grade")
    grade_label = GRADE_SLUG_TO_LABEL[grade_slug]
    if grade_label in GRADES_WITHOUT_ECA:
        raise HTTPException(400, f"{grade_label} has no ECAs")
    month = await firebase.get_month(month_id)
    if not month:
        raise HTTPException(404, "Month not found")

    entries_by_type = {}
    for eca_type in ECA_TYPE_ORDER:
        type_slug = slugify(eca_type)
        entry_id = f"{grade_slug}__{type_slug}__{month_id}"
        entries_by_type[eca_type] = await firebase.get_afterschool_entry(entry_id)

    doc_bytes = await run_in_threadpool(
        _build_eca_doc, grade_label=grade_label, month=month, entries_by_type=entries_by_type
    )
    slug = slugify(grade_label)
    filename = f'eca_{slug}_{month_id}.docx'
    return Response(
        content=doc_bytes,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )
